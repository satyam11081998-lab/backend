"""Shared helpers for building DOCX/PDF fixtures and walking an interview."""

from __future__ import annotations

import io
import uuid

from tests.conftest import CV_TEXT, JD_TEXT, auth, enable_pro, mint, run_jobs


def docx_bytes(text: str, *, table_rows=None, stamp=(2026, 1, 1, 0, 0, 0)) -> bytes:
    import docx
    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    if table_rows:
        t = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return _fixed_zip_times(buf.getvalue(), stamp)


def _fixed_zip_times(data: bytes, stamp) -> bytes:
    """python-docx stamps zip entries with the current time; pin them so the same text always
    produces the same bytes (otherwise byte-level dedupe is flaky across a second boundary)."""
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=stamp)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


def pdf_bytes(text: str) -> bytes:
    """A real text PDF via pdfium (no extra dependency)."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    pdf = pdfium.PdfDocument.new()
    width, height = 612, 792
    lines = text.splitlines() or [""]
    per_page = 45
    for start in range(0, len(lines), per_page):
        page = pdf.new_page(width, height)
        font = raw.FPDFText_LoadStandardFont(pdf.raw, b"Helvetica")
        y = height - 50
        for line in lines[start:start + per_page]:
            obj = raw.FPDFPageObj_NewTextObj(pdf.raw, b"Helvetica", 10.0)
            ws = (line + "\x00").encode("utf-16-le")
            import ctypes
            buf = ctypes.create_string_buffer(ws)
            raw.FPDFText_SetText(obj, ctypes.cast(buf, ctypes.POINTER(raw.FPDF_WCHAR)))
            raw.FPDFPageObj_Transform(obj, 1, 0, 0, 1, 50, y)
            raw.FPDFPage_InsertObject(page.raw, obj)
            y -= 15
        raw.FPDFPage_GenerateContent(page.raw)
    out = io.BytesIO()
    pdf.save(out)
    return out.getvalue()


class Candidate:
    """A user walking the product through the API."""

    def __init__(self, client, *, email="pro@example.invalid", tier="pro", uid=None, adm=False):
        self.client = client
        self.uid = uid or uuid.uuid4()
        self.email = email
        self.tier = tier
        self.adm = adm

    @property
    def h(self):
        return auth(mint(self.uid, email=self.email, tier=self.tier, adm=self.adm))

    def upload_cv(self, text=CV_TEXT, name="cv.docx"):
        r = self.client.post("/v1/documents", data={"kind": "cv"},
                             files={"file": (name, docx_bytes(text), "application/octet-stream")}, headers=self.h)
        assert r.status_code == 201, r.text
        return r.json()

    def paste_jd(self, text=JD_TEXT):
        r = self.client.post("/v1/documents/text", json={"kind": "jd", "text": text}, headers=self.h)
        assert r.status_code == 201, r.text
        return r.json()

    def create(self, cv_id, jd_id, **config):
        cfgd = {"mode": "mixed", "difficulty": "medium", "depth": "standard", "duration_minutes": 45}
        cfgd.update(config)
        return self.client.post("/v1/sessions", json={"cv_document_id": cv_id, "jd_document_id": jd_id,
                                                      "config": cfgd}, headers=self.h)

    def ready_session(self, **config):
        cv = self.upload_cv()
        jd = self.paste_jd()
        run_jobs()
        r = self.create(cv["id"], jd["id"], **config)
        assert r.status_code == 201, r.text
        sid = r.json()["id"]
        run_jobs()
        s = self.client.get(f"/v1/sessions/{sid}", headers=self.h).json()
        assert s["status"] == "ready", s
        return sid

    def turn(self, sid, content, n=[0], **kw):
        n[0] += 1
        body = {"client_turn_id": f"t{n[0]}-{uuid.uuid4().hex[:6]}", "content": content}
        body.update(kw)
        return self.client.post(f"/v1/sessions/{sid}/turns", json=body, headers=self.h)


STRONG_ANSWER = ("I led the relaunch myself because our repeat rate had dropped to 21%. I decided to cut two "
                 "under-performing SKUs and moved 30% of the budget to the lapsed-buyer segment, which meant we "
                 "gave up some reach. Within two quarters revenue grew 18% against a flat category, and I measured "
                 "it against a control region so I could attribute roughly two thirds of the lift to the change.")
WEAK_ANSWER = "We did a lot of things and the team worked hard on it."
